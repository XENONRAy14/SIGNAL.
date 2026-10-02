import io
from docx import Document
from app.resumes import extract_file,tailor

def test_extract_docx_tables():
    doc=Document();doc.add_paragraph('Expérience Python en analyse de données et automatisation.')
    table=doc.add_table(rows=1,cols=2);table.cell(0,0).text='Formation';table.cell(0,1).text='Master informatique'
    b=io.BytesIO();doc.save(b)
    result=extract_file('resume.docx',b.getvalue())
    assert 'Master informatique' in result

def test_tailoring_never_adds_missing_skills(job):
    original='Camille Martin\nDiplôme licence informatique 2024\nProjet SQL pour une association locale depuis mars 2025.'
    result,explanation=tailor(original,job)
    assert original in result
    assert all(s in original for s in explanation['highlights'])
    assert 'Python' in explanation['missing_skills']
    assert 'Python' not in result
